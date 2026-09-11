#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
照片重复性智能检测与相似检索工作台 (Photo Duplicate Scanner GUI)
版本：v2.1.0
核心功能：
1. 阶段一：海量照片多进程特征提取（四路 pHash、dHash、aHash、SHA256）与任务元数据映射关联，建立轻量 SQLite 索引库（支持断点续跑）。
2. 阶段二：基于 BK-Tree 精确汉明距离的高性能空间索引秒级候选对召回，支持全勾选控制（同任务排除、7天时间窗口、跨期留存、无条件排商户等）。
3. 成果双轨输出：
   - 轨迹 A：直接导出「相似候选清单 + 判定原因」Excel/CSV（秒级查看全局）。
   - 轨迹 B：可选同时导出「带照片离线复核包（HTML 网页 + 缩略图）」与「左右并排对比图（PAIR-xxxx.jpg）」。
4. 深度衔接：支持“先看清单，后一键将清单打包为带图复核包”的工作流。
"""

from __future__ import annotations

import os
import sys
import re
import time
import math
import json
import shutil
import random
import sqlite3
import hashlib
import zipfile
import threading
import subprocess
from datetime import datetime
from pathlib import Path
from collections import defaultdict
from typing import Dict, List, Tuple, Optional, Any, Set, Iterable
from urllib.parse import unquote, urlsplit

import numpy as np
import pandas as pd
from PIL import Image, ImageFile, ImageOps, ImageDraw
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

# 兼容 OpenCV
try:
    import cv2
except ImportError:
    cv2 = None

# Pillow 安全设置
Image.MAX_IMAGE_PIXELS = None
ImageFile.LOAD_TRUNCATED_IMAGES = True

TOOL_VERSION = "2.1.0"
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif"}
SKIP_PREFIXES = (
    ".venv",
    "review_packages",
    "相似照片对预览",
    "门头照复核_",
    "__pycache__",
)
MAPPING_COLUMNS = [
    "task_no", "task_id", "merchant_no", "merchant", "auditor_id",
    "auditor_name", "auditor_phone", "submit_time", "import_address",
    "import_lng", "import_lat", "item_id", "item_name", "doc_url",
    "upload_time", "live_address", "live_lng", "live_lat",
]
PHASH_COLUMNS = [
    ("phash_original", "原图"),
    ("phash_crop15", "去底15%"),
    ("phash_crop25", "去底25%"),
    ("phash_center90", "中心90%"),
]
HASH_COLUMNS = [col for col, _ in PHASH_COLUMNS] + ["dhash", "ahash"]
METADATA_COLUMNS = [
    "filename", "path", "task_no", "task_id", "merchant_no", "merchant",
    "auditor_id", "auditor_name", "auditor_phone", "submit_time",
    "import_address", "address_norm", "import_lng", "import_lat",
    "item_id", "item_name", "doc_url", "upload_time",
    "live_address", "live_lng", "live_lat",
]
POPCOUNT = np.array([int(i).bit_count() for i in range(256)], dtype=np.uint8)


# HTML 复核系统内嵌模板
HTML_TEMPLATE = r'''<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>门头照离线复核系统</title>
<style>
:root{--bg:#f6f6f4;--surface:#fff;--soft:#f1f0ee;--border:#e4e3e0;--text:#2c2c2b;--muted:#77736e;--blue:#2783de;--blue-soft:#e5f2fc;--red:#d9554b;--red-soft:#fce9e7;--orange:#c66a25;--orange-soft:#fbebde;--green:#2f8f62;--green-soft:#e8f1ec;--shadow:0 1px 2px rgba(0,0,0,.04),0 6px 20px rgba(0,0,0,.04)}
*{box-sizing:border-box}html,body{height:100%;margin:0}body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI","PingFang SC","Microsoft YaHei",Arial,sans-serif;background:var(--bg);color:var(--text);font-size:14px;overflow:hidden}button,input,select,textarea{font:inherit;color:inherit}button{cursor:pointer}.app{height:100%;display:grid;grid-template-rows:60px 1fr}.topbar{display:flex;align-items:center;gap:18px;background:var(--surface);border-bottom:1px solid var(--border);padding:8px 14px}.brand{display:flex;align-items:center;gap:10px;min-width:230px}.brand-mark{width:34px;height:34px;border-radius:9px;background:var(--blue-soft);color:var(--blue);display:grid;place-items:center;font-size:18px}.brand h1{font-size:17px;margin:0}.brand small{display:block;color:var(--muted);margin-top:2px}.progress{display:flex;align-items:center;gap:10px;flex:1}.track{height:7px;background:var(--soft);border-radius:99px;overflow:hidden;width:min(320px,45vw)}.bar{height:100%;background:var(--blue);width:0}.progress span{color:var(--muted);font-variant-numeric:tabular-nums}.top-actions{display:flex;gap:7px}.btn{min-height:38px;border:1px solid var(--border);border-radius:8px;background:var(--surface);padding:7px 11px}.btn:hover{background:var(--soft)}.btn.primary{background:var(--blue);border-color:var(--blue);color:#fff}.btn:focus-visible,.queue-item:focus-visible,.decision:focus-visible,input:focus-visible,select:focus-visible,textarea:focus-visible{outline:3px solid rgba(39,131,222,.22);outline-offset:1px}.layout{min-height:0;display:grid;grid-template-columns:300px minmax(590px,1fr) 330px;gap:10px;padding:10px}.panel{min-height:0;background:var(--surface);border:1px solid var(--border);border-radius:11px;box-shadow:var(--shadow);overflow:hidden}.queue{display:grid;grid-template-rows:auto 1fr}.queue-head{padding:12px;border-bottom:1px solid var(--border)}.queue-head h2{font-size:14px;margin:0 0 9px}.queue-controls{display:grid;gap:7px}.date-row{display:grid;grid-template-columns:minmax(0,1fr) minmax(0,1fr);gap:6px}.date-row input{min-width:0;padding:0 4px;font-size:12px}.filter-label{font-size:12px;color:var(--muted);margin-bottom:-3px}.queue-controls input,.queue-controls select,.reviewer input{width:100%;height:38px;border:1px solid var(--border);border-radius:8px;background:var(--surface);padding:0 9px}.reset-filter{width:100%;min-height:34px}.task-list{overflow:auto;padding:6px}.queue-item{width:100%;text-align:left;border:1px solid transparent;background:transparent;border-radius:8px;padding:9px;margin:2px 0}.queue-item:hover{background:var(--soft)}.queue-item.active{background:var(--blue-soft);border-color:#abd2f2}.queue-item.done .seq::after{content:' ✓';color:var(--green)}.queue-top{display:flex;align-items:center;justify-content:space-between;gap:6px}.seq{font-weight:700;font-size:13px}.tag{font-size:12px;border-radius:5px;padding:2px 5px}.tag.high{color:var(--red);background:var(--red-soft)}.tag.medium{color:var(--orange);background:var(--orange-soft)}.tag.low{color:var(--muted);background:var(--soft)}.task-pair{font-size:12px;color:var(--muted);line-height:1.45;margin-top:5px}.task-pair div{white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.center{display:grid;grid-template-rows:auto 1fr auto}.center-head{display:flex;align-items:center;justify-content:space-between;padding:10px 12px;border-bottom:1px solid var(--border)}.current-title{display:flex;align-items:center;gap:8px}.current-title strong{font-size:15px}.position{color:var(--muted)}.compare{min-height:0;display:grid;grid-template-columns:1fr 1fr;gap:1px;background:var(--border)}.photo-card{min-width:0;min-height:0;background:var(--surface);display:grid;grid-template-rows:auto minmax(260px,1fr) 126px}.photo-head{padding:8px 10px;border-bottom:1px solid var(--border);min-width:0}.photo-head strong{display:block;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}.photo-head span{display:block;color:var(--muted);font-size:12px;margin-top:2px}.image-frame{min-height:0;background:#171717;display:flex;align-items:center;justify-content:center;overflow:hidden;padding:6px}.image-frame img{display:block;max-width:100%;max-height:100%;width:auto;height:auto;object-fit:contain;cursor:zoom-in}.missing{color:#ddd;text-align:center}.watermark-box{height:126px;background:#111;border-top:1px solid var(--border);position:relative;overflow:hidden}.watermark-box::before{content:'底部水印放大';position:absolute;left:7px;top:6px;z-index:2;background:rgba(0,0,0,.68);color:white;border-radius:5px;padding:3px 6px;font-size:11px}.watermark-box img{width:100%;height:100%;object-fit:cover;object-position:center bottom;cursor:zoom-in}.watermark-box.hidden{display:none}.metrics{display:flex;gap:16px;align-items:center;padding:9px 12px;min-height:42px;border-top:1px solid var(--border);overflow:auto;white-space:nowrap}.metric small{color:var(--muted);margin-right:4px}.metric b{font-variant-numeric:tabular-nums}.review{display:grid;grid-template-rows:auto 1fr auto}.review-head{padding:11px 12px;border-bottom:1px solid var(--border)}.reviewer label{display:block;color:var(--muted);font-size:12px;margin-bottom:5px}.details{overflow:auto;padding:10px 12px}.detail-card{background:var(--bg);border:1px solid var(--border);border-radius:8px;padding:9px;margin-bottom:8px}.detail-card h3{font-size:13px;margin:0 0 7px}.kv{margin:6px 0}.kv small{display:block;color:var(--muted);font-size:12px;margin-bottom:1px}.kv div{line-height:1.4;word-break:break-word}.reason{background:var(--blue-soft);color:#205f94;padding:9px;border-radius:8px;line-height:1.45}.review-actions{padding:10px 12px;border-top:1px solid var(--border)}.decisions{display:grid;grid-template-columns:1fr 1fr;gap:6px}.decision{min-height:42px;border:1px solid var(--border);border-radius:8px;background:var(--surface);text-align:left;padding:7px 8px}.decision:hover{background:var(--soft)}.decision.selected{background:var(--blue-soft);border-color:var(--blue);box-shadow:inset 0 0 0 1px var(--blue)}.decision b{display:inline-block;width:18px;color:var(--muted)}.decision.full{grid-column:1/-1}.notes{width:100%;height:50px;resize:none;border:1px solid var(--border);border-radius:8px;padding:7px 9px;margin-top:7px}.nav{display:grid;grid-template-columns:1fr 1fr;gap:7px;margin-top:7px}.modal{position:fixed;inset:0;background:rgba(0,0,0,.9);z-index:20;display:none;align-items:center;justify-content:center}.modal.open{display:flex}.modal img{max-width:95vw;max-height:91vh;object-fit:contain;transform:scale(1)}.modal-tools{position:absolute;right:16px;top:14px;display:flex;gap:7px}.modal-tools .btn{background:rgba(255,255,255,.96);color:#222}.toast{position:fixed;left:50%;bottom:22px;transform:translateX(-50%);background:#222;color:#fff;border-radius:7px;padding:8px 13px;opacity:0;pointer-events:none;z-index:30}.toast.show{opacity:1}.empty{padding:24px;text-align:center;color:var(--muted)}
@media(max-width:1100px){body{overflow:auto}.app{height:auto;min-height:100%}.layout{grid-template-columns:220px minmax(540px,1fr)}.review{grid-column:1/-1;min-height:620px}.photo-card{grid-template-rows:auto 400px 120px}}
@media(max-width:760px){.topbar{height:auto;flex-wrap:wrap}.app{grid-template-rows:auto 1fr}.progress{order:3;width:100%}.layout{display:block}.queue{height:300px;margin-bottom:10px}.center{height:auto;min-height:0;grid-template-rows:auto auto auto;margin-bottom:10px}.compare{grid-template-columns:1fr}.photo-card{grid-template-rows:auto 300px 110px}.review{min-height:760px}.top-actions{width:100%;overflow:auto}.track{width:100%}}
@media(prefers-color-scheme:dark){:root{--bg:#191919;--surface:#202020;--soft:#383836;--border:rgba(255,255,255,.18);--text:#fff;--muted:rgba(255,255,255,.65);--blue:#5e9fe8;--blue-soft:rgba(94,159,232,.14);--red:#e97366;--red-soft:rgba(233,115,102,.14);--orange:#de9255;--orange-soft:rgba(222,146,85,.14);--green:#72bc8f;--green-soft:rgba(114,188,143,.14);--shadow:none}.reason{color:#a8d2f5}}
</style>
</head>
<body>
<div class="app">
<header class="topbar">
  <div class="brand"><div class="brand-mark">◎</div><div><h1>门头照复核</h1><small id="batchLabel"></small></div></div>
  <div class="progress"><div class="track"><div class="bar" id="progressBar"></div></div><span id="progressText"></span></div>
  <div class="top-actions"><button class="btn" id="watermarkBtn">隐藏水印放大</button><button class="btn" id="importBtn">导入进度</button><button class="btn" id="backupBtn">备份 JSON</button><button class="btn primary" id="exportBtn">导出 CSV</button><input type="file" id="importFile" accept="application/json" hidden></div>
</header>
<div class="layout">
  <aside class="queue panel"><div class="queue-head"><h2>任务号列表</h2><div class="queue-controls"><input id="taskSearch" placeholder="搜索任务号"><select id="operatorFilter"><option value="">全部作业员</option></select><input id="addressSearch" placeholder="筛选地址关键词"><div class="filter-label">提交时间（任一照片命中）</div><div class="date-row"><input type="date" id="startDate" aria-label="开始日期"><input type="date" id="endDate" aria-label="结束日期"></div><select id="filter"><option>全部</option><option>未复核</option><option>高优先级</option><option>中优先级</option><option>已确认复用</option><option>同店正常</option></select><button class="btn reset-filter" id="resetFilters">重置筛选</button></div></div><div class="task-list" id="taskList"></div></aside>
  <main class="center panel"><div class="center-head"><div class="current-title"><strong id="groupId"></strong><span class="tag" id="priority"></span><span class="position" id="positionText"></span></div><span class="position">点击图片可放大</span></div>
    <section class="compare">
      <article class="photo-card"><div class="photo-head"><strong id="file1"></strong><span id="taskHead1"></span></div><div class="image-frame"><img id="img1" alt="候选照片1"><div class="missing" id="missing1" hidden>照片缺失</div></div><div class="watermark-box" id="watermarkBox1"><img id="watermark1" alt="照片1底部水印"></div></article>
      <article class="photo-card"><div class="photo-head"><strong id="file2"></strong><span id="taskHead2"></span></div><div class="image-frame"><img id="img2" alt="候选照片2"><div class="missing" id="missing2" hidden>照片缺失</div></div><div class="watermark-box" id="watermarkBox2"><img id="watermark2" alt="照片2底部水印"></div></article>
    </section><div class="metrics" id="metrics"></div>
  </main>
  <aside class="review panel"><div class="review-head"><div class="reviewer"><label for="reviewer">复核人</label><input id="reviewer" placeholder="请输入姓名"></div></div><div class="details"><div id="detail1" class="detail-card"></div><div id="detail2" class="detail-card"></div><div class="reason" id="reason"></div></div><div class="review-actions"><div class="decisions"><button class="decision" data-value="确认复用"><b>1</b>确认复用</button><button class="decision" data-value="正常不同照片"><b>2</b>正常不同</button><button class="decision" data-value="同店正常拍摄"><b>3</b>同店正常</button><button class="decision" data-value="系统图"><b>4</b>系统图</button><button class="decision full" data-value="无法判断"><b>5</b>无法判断</button></div><textarea id="notes" class="notes" placeholder="可选备注"></textarea><div class="nav"><button class="btn" id="prevBtn">← 上一组</button><button class="btn primary" id="nextBtn">下一组 →</button></div></div></aside>
</div>
</div>
<div class="modal" id="modal"><div class="modal-tools"><button class="btn" id="zoomOut">−</button><button class="btn" id="zoomReset">100%</button><button class="btn" id="zoomIn">＋</button><button class="btn" id="closeModal">关闭</button></div><img id="modalImage" alt="大图预览"></div><div class="toast" id="toast"></div>
<script>
const BATCH_ID=__BATCH_ID_JSON__;const ITEMS=__CANDIDATES_JSON__;const VERDICTS=['确认复用','正常不同照片','同店正常拍摄','系统图','无法判断'];const storageKey='photo-review-v2:'+BATCH_ID;
let state={reviewer:'',reviews:{},currentId:ITEMS[0]?.id||'',filter:'全部',taskSearch:'',operator:'',addressSearch:'',startDate:'',endDate:'',showWatermark:true};try{const saved=JSON.parse(localStorage.getItem(storageKey)||'null');if(saved)state={...state,...saved};if(state.search&&!state.taskSearch)state.taskSearch=state.search}catch(e){}
let visible=[];let zoom=1;const $=id=>document.getElementById(id);function safe(v){return v===null||v===undefined||v===''?'—':String(v)}function esc(v){return safe(v).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}function reviewFor(item){return state.reviews[item.id]||{verdict:'',notes:'',reviewer:'',reviewTime:''}}function save(){localStorage.setItem(storageKey,JSON.stringify(state))}
function dateKey(v){const m=String(v||'').match(/(\d{4})[\/-](\d{1,2})[\/-](\d{1,2})/);return m?`${m[1]}-${String(m[2]).padStart(2,'0')}-${String(m[3]).padStart(2,'0')}`:''}function renderOperatorOptions(){const values=[...new Set(ITEMS.flatMap(i=>[i.auditor1,i.auditor2]).filter(Boolean))].sort((a,b)=>String(a).localeCompare(String(b),'zh-CN'));$('operatorFilter').innerHTML='<option value="">全部作业员</option>'+values.map(v=>`<option value="${esc(v)}">${esc(v)}</option>`).join('');$('operatorFilter').value=state.operator}function rebuildVisible(){const tq=state.taskSearch.trim().toLowerCase();const aq=state.addressSearch.trim().toLowerCase();visible=ITEMS.filter(item=>{const r=reviewFor(item);let ok=true;if(state.filter==='未复核')ok=!r.verdict;else if(state.filter==='高优先级')ok=item.priority==='高';else if(state.filter==='中优先级')ok=item.priority==='中';else if(state.filter==='已确认复用')ok=r.verdict==='确认复用';else if(state.filter==='同店正常')ok=r.verdict==='同店正常拍摄';if(!ok)return false;if(state.operator&&![item.auditor1,item.auditor2].includes(state.operator))return false;if(tq&&![item.task1,item.task2,item.file1,item.file2,item.group].join(' ').toLowerCase().includes(tq))return false;if(aq&&![item.address1,item.address2].join(' ').toLowerCase().includes(aq))return false;if(state.startDate||state.endDate){const dates=[dateKey(item.time1),dateKey(item.time2)].filter(Boolean);const dateOk=dates.some(d=>(!state.startDate||d>=state.startDate)&&(!state.endDate||d<=state.endDate));if(!dateOk)return false}return true});if(!visible.some(x=>x.id===state.currentId))state.currentId=visible[0]?.id||''}
function current(){return ITEMS.find(x=>x.id===state.currentId)||null}function priorityClass(v){return v==='高'?'high':v==='中'?'medium':'low'}
function renderList(){const html=visible.map((item,index)=>{const r=reviewFor(item);return`<button class="queue-item ${item.id===state.currentId?'active':''} ${r.verdict?'done':''}" data-id="${esc(item.id)}"><div class="queue-top"><span class="seq">${index+1}. ${esc(item.group)}</span><span class="tag ${priorityClass(item.priority)}">${esc(item.priority)}</span></div><div class="task-pair"><div>${esc(item.task1)}</div><div>↔ ${esc(item.task2)}</div></div></button>`}).join('');$('taskList').innerHTML=html||'<div class="empty">没有符合条件的任务</div>';document.querySelectorAll('.queue-item').forEach(btn=>btn.onclick=()=>{state.currentId=btn.dataset.id;render()});const active=document.querySelector('.queue-item.active');if(active)active.scrollIntoView({block:'nearest'})}
function detailHtml(title,item,side){return`<h3>${title}</h3><div class="kv"><small>任务号</small><div>${esc(item['task'+side])}</div></div><div class="kv"><small>导入地址</small><div>${esc(item['address'+side])}</div></div><div class="kv"><small>采集项</small><div>${esc(item['item'+side])}</div></div><div class="kv"><small>提交时间</small><div>${esc(item['time'+side])}</div></div><div class="kv"><small>作业员</small><div>${esc(item['auditor'+side])}</div></div>`}function metric(label,value){return`<div class="metric"><small>${label}</small><b>${esc(value)}</b></div>`}
function setImage(img,missing,src){img.hidden=false;missing.hidden=true;img.onerror=()=>{img.hidden=true;missing.hidden=false};img.src=src}
function render(){rebuildVisible();renderList();renderOperatorOptions();$('batchLabel').textContent=BATCH_ID;$('reviewer').value=state.reviewer;$('filter').value=state.filter;$('taskSearch').value=state.taskSearch;$('addressSearch').value=state.addressSearch;$('startDate').value=state.startDate;$('endDate').value=state.endDate;const reviewed=ITEMS.filter(x=>reviewFor(x).verdict).length;$('progressText').textContent=`已复核 ${reviewed} / ${ITEMS.length}`;$('progressBar').style.width=(ITEMS.length?reviewed/ITEMS.length*100:0)+'%';const item=current();if(!item){$('groupId').textContent='暂无任务';return}const r=reviewFor(item);const pos=visible.findIndex(x=>x.id===item.id);$('groupId').textContent=item.group;$('priority').textContent=item.priority+'优先级';$('priority').className='tag '+priorityClass(item.priority);$('positionText').textContent=`${pos+1} / ${visible.length}`;$('file1').textContent=item.file1;$('file2').textContent=item.file2;$('taskHead1').textContent=item.task1;$('taskHead2').textContent=item.task2;setImage($('img1'),$('missing1'),item.image1);setImage($('img2'),$('missing2'),item.image2);$('watermark1').src=item.image1;$('watermark2').src=item.image2;$('watermarkBox1').classList.toggle('hidden',!state.showWatermark);$('watermarkBox2').classList.toggle('hidden',!state.showWatermark);$('watermarkBtn').textContent=state.showWatermark?'隐藏水印放大':'显示水印放大';$('detail1').innerHTML=detailHtml('照片 1',item,'1');$('detail2').innerHTML=detailHtml('照片 2',item,'2');$('reason').textContent=item.reason;$('metrics').innerHTML=metric('pHash',item.phash)+metric('最佳区域',item.phashArea)+metric('dHash',item.dhash)+metric('ORB内点',item.orbInliers)+metric('内点比例',item.orbRatio);$('notes').value=r.notes||'';document.querySelectorAll('.decision').forEach(b=>b.classList.toggle('selected',b.dataset.value===r.verdict));save();preloadNext()}
function move(delta){if(!visible.length)return;let i=visible.findIndex(x=>x.id===state.currentId);i=Math.min(Math.max(i+delta,0),visible.length-1);state.currentId=visible[i].id;render()}function setVerdict(value){const item=current();if(!item)return;state.reviews[item.id]={...reviewFor(item),verdict:value,notes:$('notes').value,reviewer:state.reviewer,reviewTime:new Date().toISOString()};save();toast('已记录：'+value);move(1)}function preloadNext(){const i=visible.findIndex(x=>x.id===state.currentId);const n=visible[i+1];if(n){new Image().src=n.image1;new Image().src=n.image2}}function toast(t){$('toast').textContent=t;$('toast').classList.add('show');setTimeout(()=>$('toast').classList.remove('show'),1200)}
function csvCell(v){const s=v==null?'':String(v);return'"'+s.replaceAll('"','""')+'"'}function download(name,text,type){const blob=new Blob([text],{type});const url=URL.createObjectURL(blob);const a=document.createElement('a');a.href=url;a.download=name;a.click();setTimeout(()=>URL.revokeObjectURL(url),800)}function exportCsv(){const h=['batch_id','candidate_id','group','priority','file1','file2','task1','task2','address1','address2','verdict','notes','reviewer','review_time','phash','orb_inliers','orb_ratio'];const rows=[h.join(',')];ITEMS.forEach(item=>{const r=reviewFor(item);rows.push([BATCH_ID,item.id,item.group,item.priority,item.file1,item.file2,item.task1,item.task2,item.address1,item.address2,r.verdict,r.notes,r.reviewer,r.reviewTime,item.phash,item.orbInliers,item.orbRatio].map(csvCell).join(','))});download(`${BATCH_ID}_复核结果.csv`,'\ufeff'+rows.join('\r\n'),'text/csv;charset=utf-8')}function backup(){download(`${BATCH_ID}_进度备份.json`,JSON.stringify({batchId:BATCH_ID,state},null,2),'application/json')}function importBackup(file){const reader=new FileReader();reader.onload=()=>{try{const data=JSON.parse(reader.result);if(data.batchId!==BATCH_ID)throw new Error('批次不一致');state={...state,...data.state};save();render();toast('进度已恢复')}catch(e){alert('导入失败：'+e.message)}};reader.readAsText(file)}
function openModal(src){zoom=1;$('modalImage').src=src;$('modalImage').style.transform='scale(1)';$('zoomReset').textContent='100%';$('modal').classList.add('open')}function changeZoom(d){zoom=Math.min(Math.max(zoom+d,.5),4);$('modalImage').style.transform=`scale(${zoom})`;$('zoomReset').textContent=Math.round(zoom*100)+'%'}
$('reviewer').onchange=e=>{state.reviewer=e.target.value.trim();save()};$('taskSearch').oninput=e=>{state.taskSearch=e.target.value;render()};$('operatorFilter').onchange=e=>{state.operator=e.target.value;render()};$('addressSearch').oninput=e=>{state.addressSearch=e.target.value;render()};$('startDate').onchange=e=>{state.startDate=e.target.value;render()};$('endDate').onchange=e=>{state.endDate=e.target.value;render()};$('filter').onchange=e=>{state.filter=e.target.value;render()};$('resetFilters').onclick=()=>{state.filter='全部';state.taskSearch='';state.operator='';state.addressSearch='';state.startDate='';state.endDate='';render()};$('watermarkBtn').onclick=()=>{state.showWatermark=!state.showWatermark;render()};$('notes').onchange=()=>{const item=current();if(item){state.reviews[item.id]={...reviewFor(item),notes:$('notes').value,reviewer:state.reviewer};save()}};document.querySelectorAll('.decision').forEach(b=>b.onclick=()=>setVerdict(b.dataset.value));$('prevBtn').onclick=()=>move(-1);$('nextBtn').onclick=()=>move(1);$('exportBtn').onclick=exportCsv;$('backupBtn').onclick=backup;$('importBtn').onclick=()=>$('importFile').click();$('importFile').onchange=e=>e.target.files[0]&&importBackup(e.target.files[0]);['img1','img2','watermark1','watermark2'].forEach(id=>$(id).onclick=()=>openModal($(id).src));$('closeModal').onclick=()=>$('modal').classList.remove('open');$('modal').onclick=e=>{if(e.target===$('modal'))$('modal').classList.remove('open')};$('zoomIn').onclick=()=>changeZoom(.25);$('zoomOut').onclick=()=>changeZoom(-.25);$('zoomReset').onclick=()=>{zoom=1;$('modalImage').style.transform='scale(1)';$('zoomReset').textContent='100%'};document.addEventListener('keydown',e=>{if(['INPUT','TEXTAREA','SELECT'].includes(document.activeElement.tagName))return;if(e.key>='1'&&e.key<='5'){setVerdict(VERDICTS[Number(e.key)-1]);e.preventDefault()}else if(e.key==='ArrowLeft'){move(-1);e.preventDefault()}else if(e.key==='ArrowRight'){move(1);e.preventDefault()}else if(e.key==='Escape')$('modal').classList.remove('open')});render();
</script>
</body>
</html>'''

README_TEMPLATE = '''门头照离线复核包使用说明

1. 先完整解压 ZIP（若未打包为 ZIP，直接在文件夹中打开）。
2. 双击“开始复核.html”，在 Chrome 或 Edge 中打开。完全离线无依赖。
3. 键盘快捷键：1确认复用、2正常不同、3同店正常、4系统图、5无法判断；左右方向键切换上一对/下一对。
4. 审核完成后，点击右上角“导出 CSV”获取复核结果。
'''


# ==============================================================================
# 算法核心工具函数
# ==============================================================================

def dct_matrix(size: int = 32) -> np.ndarray:
    x = np.arange(size)
    u = np.arange(size)[:, None]
    matrix = np.cos((2 * x + 1) * u * np.pi / (2 * size))
    matrix[0, :] *= np.sqrt(1 / size)
    matrix[1:, :] *= np.sqrt(2 / size)
    return matrix.astype(np.float32)

DCT32 = dct_matrix(32)

def gray_array(image: Image.Image, size: tuple[int, int]) -> np.ndarray:
    resized = image.convert("L").resize(size, Image.Resampling.LANCZOS)
    return np.asarray(resized, dtype=np.float32)

def phash_hex(image: Image.Image) -> str:
    pixels = gray_array(image, (32, 32))
    transformed = DCT32 @ pixels @ DCT32.T
    low = transformed[:8, :8]
    threshold = float(np.median(low.reshape(-1)[1:]))
    bits = (low > threshold).reshape(-1)
    val = 0
    for bit in bits:
        val = (val << 1) | int(bit)
    return f"{val:016x}"

def dhash_hex(image: Image.Image) -> str:
    pixels = gray_array(image, (9, 8))
    bits = (pixels[:, 1:] > pixels[:, :-1]).reshape(-1)
    val = 0
    for bit in bits:
        val = (val << 1) | int(bit)
    return f"{val:016x}"

def ahash_hex(image: Image.Image) -> str:
    pixels = gray_array(image, (8, 8))
    bits = (pixels > float(pixels.mean())).reshape(-1)
    val = 0
    for bit in bits:
        val = (val << 1) | int(bit)
    return f"{val:016x}"

def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "null", "\\n"} else text

def normalize_address(value: Any) -> str:
    text = clean_text(value).lower()
    return re.sub(r"[\s,，。;；:：、（）()\-]+", "", text)

def normalize_merchant(value: Any) -> str:
    text = clean_text(value)
    if not text or text.startswith(("http://", "https://")):
        return ""
    if re.fullmatch(r"[+\-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+\-]?\d+)?", text):
        return ""
    if text.startswith("@encrypt@"):
        return text
    return re.sub(r"[\s,，。;；:：、（）()\-_]+", "", text).lower()

def filename_from_url(value: Any) -> str:
    text = clean_text(value)
    if not text.startswith(("http://", "https://")):
        return ""
    try:
        return Path(unquote(urlsplit(text).path)).name
    except Exception:
        return ""

def parse_submit_time(value: Any) -> float:
    text = clean_text(value)
    if not text:
        return float("nan")
    try:
        number = float(text)
    except (TypeError, ValueError):
        number = None
    if number is not None and math.isfinite(number):
        if number > 1e14:
            return number / 1_000_000.0
        if number > 1e11:
            return number / 1_000.0
        if number > 1e9:
            return number
        if 20_000 <= number <= 80_000:
            return (pd.Timestamp("1899-12-30") + pd.to_timedelta(number, unit="D")).timestamp()
    parsed = pd.to_datetime(text, errors="coerce")
    if pd.isna(parsed):
        return float("nan")
    return float(parsed.timestamp())

def parse_submit_times(records: pd.DataFrame) -> np.ndarray:
    source = records.get("submit_time", pd.Series([""] * len(records)))
    return np.fromiter((parse_submit_time(v) for v in source), dtype=np.float64, count=len(records))

def parse_hash(value: Any) -> int:
    text = clean_text(value)
    if not text:
        raise ValueError("空哈希")
    return int(text, 16)

def hamming(first: int, second: int) -> int:
    return (first ^ second).bit_count()

def hamming_array(values: np.ndarray, left: np.ndarray, right: np.ndarray, chunk_size: int = 200000) -> np.ndarray:
    output = np.empty(len(left), dtype=np.uint8)
    for start in range(0, len(left), chunk_size):
        end = min(start + chunk_size, len(left))
        xor_values = np.bitwise_xor(values[left[start:end]], values[right[start:end]])
        byte_view = xor_values.view(np.uint8).reshape(-1, 8)
        output[start:end] = POPCOUNT[byte_view].sum(axis=1, dtype=np.uint16)
    return output


# ==============================================================================
# BK-Tree 空间索引树
# ==============================================================================

class BKTree:
    def __init__(self, values: Iterable[int]):
        values = list(values)
        if not values:
            self.root = None
            return
        random.Random(20260822).shuffle(values)
        self.root = [values[0], {}]
        for v in values[1:]:
            self.add(v)

    def add(self, value: int) -> None:
        node = self.root
        while True:
            distance = hamming(value, node[0])
            child = node[1].get(distance)
            if child is None:
                node[1][distance] = [value, {}]
                return
            node = child

    def search(self, query: int, radius: int) -> List[int]:
        if self.root is None:
            return []
        found: List[int] = []
        stack = [self.root]
        while stack:
            value, children = stack.pop()
            distance = hamming(query, value)
            if distance <= radius:
                found.append(value)
            low, high = distance - radius, distance + radius
            for edge, child in children.items():
                if low <= edge <= high:
                    stack.append(child)
        return found


# ==============================================================================
# 多进程单张提取 Worker
# ==============================================================================

def extract_one_photo(path_str: str) -> dict:
    path = Path(path_str)
    started = time.time()
    try:
        stat = path.stat()
        file_sha = sha256_file(path)
        with Image.open(path) as source:
            image = ImageOps.exif_transpose(source).convert("RGB")
        width, height = image.size
        pixel_digest = hashlib.sha256()
        pixel_digest.update(f"{width}x{height}|RGB|".encode("ascii"))
        pixel_digest.update(image.tobytes())

        crop15 = image.crop((0, 0, width, max(1, int(height * 0.85))))
        crop25 = image.crop((0, 0, width, max(1, int(height * 0.75))))
        left, top = int(width * 0.05), int(height * 0.05)
        right, bottom = max(left + 1, int(width * 0.95)), max(top + 1, int(height * 0.95))
        center90 = image.crop((left, top, right, bottom))

        return {
            "filename": path.name,
            "path": str(path),
            "size": stat.st_size,
            "mtime_ns": stat.st_mtime_ns,
            "width": width,
            "height": height,
            "sha256": file_sha,
            "pixel_sha256": pixel_digest.hexdigest(),
            "phash_original": phash_hex(image),
            "phash_crop15": phash_hex(crop15),
            "phash_crop25": phash_hex(crop25),
            "phash_center90": phash_hex(center90),
            "dhash": dhash_hex(image),
            "ahash": ahash_hex(image),
            "status": "ok",
            "error": "",
            "elapsed_ms": int((time.time() - started) * 1000),
        }
    except Exception as exc:
        try:
            stat = path.stat()
            size, mtime_ns = stat.st_size, stat.st_mtime_ns
        except Exception:
            size, mtime_ns = 0, 0
        return {
            "filename": path.name,
            "path": str(path),
            "size": size,
            "mtime_ns": mtime_ns,
            "status": "failed",
            "error": f"{type(exc).__name__}: {exc}",
            "elapsed_ms": int((time.time() - started) * 1000),
        }


# ==============================================================================
# 核心查重与复核打包引擎 (DuplicateScanEngine)
# ==============================================================================

class DuplicateScanEngine:
    def __init__(self, log_callback=None, progress_callback=None, status_callback=None):
        self.log_callback = log_callback or (lambda msg, level="INFO": print(f"[{level}] {msg}"))
        self.progress_callback = progress_callback or (lambda cur, total: None)
        self.status_callback = status_callback or (lambda text: None)
        self.is_cancelled = False

    def cancel(self):
        self.is_cancelled = True

    # ---------------- 阶段一：建立/增量更新 SQLite 特征库 ----------------
    def build_photo_index(
        self,
        root_dir: Path,
        mapping_path: Optional[Path],
        db_path: Path,
        workers: int = 4,
        commit_every: int = 200,
        refresh: bool = False,
    ) -> dict:
        self.status_callback("正在扫描照片文件与读取映射表...")
        self.log_callback(f"照片根目录：{root_dir}")
        self.log_callback(f"特征库目标：{db_path}")

        mapping_lookup = {}
        if mapping_path and mapping_path.is_file():
            self.log_callback(f"读取映射表：{mapping_path.name}")
            try:
                mapping_lookup, _ = self._load_mapping_file(mapping_path)
                self.log_callback(f"映射表加载完成：解析到 {len(mapping_lookup):,} 条有效记录")
            except Exception as e:
                self.log_callback(f"映射表读取警告：{e}（将继续处理未绑定元数据的照片）", "WARN")

        self.log_callback("正在遍历目录收集图片文件...")
        image_paths = []
        for path in root_dir.rglob("*"):
            if self.is_cancelled:
                return {"status": "cancelled"}
            if not path.is_file() or path.name.startswith("._"):
                continue
            if path.suffix.lower() not in IMAGE_EXTENSIONS:
                continue
            if any(part.startswith(SKIP_PREFIXES) for part in path.parts):
                continue
            image_paths.append(path)

        total_files = len(image_paths)
        self.log_callback(f"共发现 {total_files:,} 张符合条件的图片")
        if total_files == 0:
            part_files = list(root_dir.glob("*.part"))
            if part_files:
                raise ValueError(
                    f"所选目录中仅发现 {len(part_files)} 个未下载完成的临时文件（.part），未找到已完成的图片！\n\n"
                    f"请检查是否选错了文件夹，真正已完成下载的照片目录位于：\n"
                    f"/Users/jun/Downloads/照片重复性检测/pos-jianhang2_已合并_photos （内含 30,000+ 张照片）"
                )
            raise ValueError(f"在目录中未找到任何常见格式照片（.jpg, .png 等）：\n{root_dir}")

        conn = self._init_sqlite(db_path)
        cached = {}
        if not refresh:
            rows = conn.execute("SELECT filename, size, mtime_ns, status FROM photos").fetchall()
            cached = {r[0]: (r[1], r[2], r[3]) for r in rows}
            self.log_callback(f"检测到已有 SQLite 缓存：{len(cached):,} 条记录")

        to_process = []
        for p in image_paths:
            fname = p.name
            stat = p.stat()
            if fname in cached and not refresh:
                c_size, c_mtime, c_status = cached[fname]
                if c_status == "ok" and c_size == stat.st_size and c_mtime == stat.st_mtime_ns:
                    continue
            to_process.append(str(p))

        self.log_callback(f"本次需实际提取特征的照片数：{len(to_process):,} 张（已跳过缓存一致项）")
        if not to_process:
            self.status_callback("特征库已为最新，无需重算")
            self.log_callback("所有照片均已在特征库中且无修改，阶段一完成！", "SUCCESS")
            return {"status": "ok", "total": total_files, "processed": 0}

        from concurrent.futures import ProcessPoolExecutor, as_completed

        processed_count = 0
        success_count = 0
        pending_records = []
        start_time = time.time()
        max_w = max(1, min(workers, os.cpu_count() or 4))
        self.status_callback(f"正在多进程并发提取特征（{max_w} 核心）...")
        self.log_callback(f"启动多进程特征提取引擎，并发数：{max_w}")

        with ProcessPoolExecutor(max_workers=max_w) as executor:
            future_to_file = {executor.submit(extract_one_photo, p): p for p in to_process}
            for future in as_completed(future_to_file):
                if self.is_cancelled:
                    executor.shutdown(wait=False, cancel_futures=True)
                    self.log_callback("用户取消了特征建库任务", "WARN")
                    return {"status": "cancelled"}

                res = future.result()
                fname = res.get("filename", "")
                m_info = mapping_lookup.get(fname, {})
                pending_records.append((res, m_info))
                processed_count += 1
                if res.get("status") == "ok":
                    success_count += 1

                if len(pending_records) >= commit_every:
                    self._commit_batch(conn, pending_records)
                    pending_records.clear()
                    conn.commit()

                if processed_count % 50 == 0 or processed_count == len(to_process):
                    elapsed = time.time() - start_time
                    speed = processed_count / max(0.1, elapsed)
                    remain = (len(to_process) - processed_count) / max(0.1, speed)
                    self.progress_callback(processed_count, len(to_process))
                    self.status_callback(
                        f"特征建库中：{processed_count:,}/{len(to_process):,} ({speed:.1f}张/秒) 剩余约{remain:.0f}秒"
                    )

        if pending_records:
            self._commit_batch(conn, pending_records)
            pending_records.clear()
            conn.commit()

        conn.close()
        self.log_callback(f"阶段一建库完成！成功写入 {success_count:,} 张照片特征", "SUCCESS")
        return {"status": "ok", "total": total_files, "processed": processed_count, "success": success_count}

    def _load_mapping_file(self, path: Path) -> Tuple[Dict[str, dict], dict]:
        data = None
        if path.suffix.lower() == ".csv":
            for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
                try:
                    data = pd.read_csv(path, dtype=str, low_memory=False, encoding=enc)
                    break
                except Exception:
                    continue
        else:
            data = pd.read_excel(path, dtype=str)

        if data is None:
            raise RuntimeError(f"无法读取表格：{path}")

        raw_rows = len(data)
        resolved_url = pd.Series("", index=data.index, dtype="object")
        if "doc_url" in data.columns:
            vals = data["doc_url"].fillna("").astype(str)
            mask = vals.str.startswith(("http://", "https://"))
            resolved_url.loc[mask] = vals.loc[mask]

        for col in data.columns:
            missing = resolved_url.eq("")
            if not missing.any():
                break
            vals = data[col].fillna("").astype(str)
            mask = missing & vals.str.startswith(("http://", "https://"))
            resolved_url.loc[mask] = vals.loc[mask]

        data = data.assign(_resolved_url=resolved_url)
        data = data[data["_resolved_url"] != ""].copy()
        data["_filename"] = data["_resolved_url"].map(filename_from_url)
        data = data[data["_filename"] != ""].copy()
        data["doc_url"] = data["_resolved_url"]

        if "import_address" in data.columns:
            data["address_norm"] = data["import_address"].map(normalize_address)
        else:
            data["address_norm"] = ""

        data = data.drop_duplicates("_filename", keep="first")
        avail = [c for c in MAPPING_COLUMNS if c in data.columns]
        lookup = {}
        for row in data[["_filename", "address_norm", *avail]].to_dict(orient="records"):
            fname = row.pop("_filename")
            lookup[fname] = {k: clean_text(v) for k, v in row.items()}

        return lookup, {"raw_rows": raw_rows, "valid": len(lookup)}

    def _init_sqlite(self, db_path: Path) -> sqlite3.Connection:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(db_path)
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA synchronous=NORMAL")
        conn.execute("PRAGMA temp_store=MEMORY")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS photos (
                filename TEXT PRIMARY KEY, path TEXT NOT NULL, size INTEGER, mtime_ns INTEGER,
                width INTEGER, height INTEGER, sha256 TEXT, pixel_sha256 TEXT,
                phash_original TEXT, phash_crop15 TEXT, phash_crop25 TEXT, phash_center90 TEXT,
                dhash TEXT, ahash TEXT, task_no TEXT, task_id TEXT, merchant_no TEXT,
                merchant TEXT, auditor_id TEXT, auditor_name TEXT, auditor_phone TEXT,
                submit_time TEXT, import_address TEXT, address_norm TEXT, import_lng TEXT,
                import_lat TEXT, item_id TEXT, item_name TEXT, doc_url TEXT, upload_time TEXT,
                live_address TEXT, live_lng TEXT, live_lat TEXT, status TEXT NOT NULL,
                error TEXT, elapsed_ms INTEGER, processed_at TEXT, tool_version TEXT
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_sha256 ON photos(sha256)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_pixel_sha256 ON photos(pixel_sha256)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_phash_original ON photos(phash_original)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_task_no ON photos(task_no)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_address_norm ON photos(address_norm)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_photos_status ON photos(status)")
        conn.commit()
        return conn

    def _commit_batch(self, conn: sqlite3.Connection, batch: List[Tuple[dict, dict]]):
        now = datetime.now().isoformat(timespec="seconds")
        items = []
        for res, mapping in batch:
            m = {col: "" for col in MAPPING_COLUMNS}
            m.update(mapping or {})
            m["address_norm"] = normalize_address(m.get("import_address", ""))
            row = (
                res.get("filename", ""), res.get("path", ""), res.get("size", 0), res.get("mtime_ns", 0),
                res.get("width"), res.get("height"), res.get("sha256"), res.get("pixel_sha256"),
                res.get("phash_original"), res.get("phash_crop15"), res.get("phash_crop25"), res.get("phash_center90"),
                res.get("dhash"), res.get("ahash"), m.get("task_no", ""), m.get("task_id", ""),
                m.get("merchant_no", ""), m.get("merchant", ""), m.get("auditor_id", ""), m.get("auditor_name", ""),
                m.get("auditor_phone", ""), m.get("submit_time", ""), m.get("import_address", ""), m.get("address_norm", ""),
                m.get("import_lng", ""), m.get("import_lat", ""), m.get("item_id", ""), m.get("item_name", ""),
                m.get("doc_url", ""), m.get("upload_time", ""), m.get("live_address", ""), m.get("live_lng", ""),
                m.get("live_lat", ""), res.get("status", "failed"), res.get("error", ""), res.get("elapsed_ms", 0),
                now, TOOL_VERSION,
            )
            items.append(row)

        conn.executemany(
            """
            INSERT INTO photos (
                filename, path, size, mtime_ns, width, height, sha256, pixel_sha256,
                phash_original, phash_crop15, phash_crop25, phash_center90, dhash, ahash,
                task_no, task_id, merchant_no, merchant, auditor_id, auditor_name, auditor_phone,
                submit_time, import_address, address_norm, import_lng, import_lat, item_id,
                item_name, doc_url, upload_time, live_address, live_lng, live_lat,
                status, error, elapsed_ms, processed_at, tool_version
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            ON CONFLICT(filename) DO UPDATE SET
                path=excluded.path, size=excluded.size, mtime_ns=excluded.mtime_ns,
                width=excluded.width, height=excluded.height, sha256=excluded.sha256,
                pixel_sha256=excluded.pixel_sha256, phash_original=excluded.phash_original,
                phash_crop15=excluded.phash_crop15, phash_crop25=excluded.phash_crop25,
                phash_center90=excluded.phash_center90, dhash=excluded.dhash, ahash=excluded.ahash,
                status=excluded.status, error=excluded.error, elapsed_ms=excluded.elapsed_ms,
                processed_at=excluded.processed_at
            """,
            items,
        )

    # ---------------- 阶段二：相似检索与双轨成果输出 ----------------
    def scan_similarity_v3(
        self,
        db_path: Path,
        output_excel_path: Path,
        output_csv_path: Path,
        phash_threshold: int = 8,
        dhash_threshold: int = 8,
        # 业务过滤勾选项
        filter_same_task: bool = True,
        use_time_window: bool = True,
        same_entity_window_days: float = 7.0,
        exclude_all_same_merchant: bool = False,
        # 成果勾选项
        export_list: bool = True,
        export_review_package: bool = False,
        export_pair_previews: bool = False,
        photo_root_dir: Optional[Path] = None,
        max_candidates: int = 2000000,
        excel_limit: int = 200000,
        dense_group_size: int = 20,
    ) -> dict:
        self.status_callback("正在读取 SQLite 特征库...")
        self.log_callback(f"连接特征库：{db_path.name}")

        if not db_path.is_file():
            raise FileNotFoundError(f"特征库文件不存在：{db_path}")

        conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
        try:
            avail = {r[1] for r in conn.execute("PRAGMA table_info(photos)")}
            req = {"filename", "path", "status", *HASH_COLUMNS}
            missing = req - avail
            if missing:
                raise ValueError(f"SQLite photos 表缺少关键列：{sorted(missing)}")
            sel_cols = [c for c in METADATA_COLUMNS + HASH_COLUMNS if c in avail]
            quoted = ", ".join(f'"{c}"' for c in sel_cols)
            records = pd.read_sql_query(
                f"SELECT {quoted} FROM photos WHERE status='ok' ORDER BY filename",
                conn,
                dtype=str,
            )
        finally:
            conn.close()

        for c in HASH_COLUMNS:
            records = records[records[c].fillna("").astype(str).str.fullmatch(r"[0-9a-fA-F]{16}")]
        records = records.reset_index(drop=True)
        total_photos = len(records)
        self.log_callback(f"成功加载有效照片记录：{total_photos:,} 条")
        if total_photos < 2:
            raise ValueError("特征库中有完整哈希的有效照片不足 2 张，无法进行相似比对！")

        hashes = {}
        for c in HASH_COLUMNS:
            hashes[c] = np.fromiter((parse_hash(v) for v in records[c]), dtype=np.uint64, count=total_photos)

        # BK-Tree 空间检索
        self.status_callback("正在通过 BK-Tree 空间树秒级召回相似对...")
        self.log_callback(
            f"检索配置：pHash={phash_threshold}, dHash={dhash_threshold}, "
            f"同任务过滤={filter_same_task}, 时间窗口={same_entity_window_days if use_time_window else '关闭'}, "
            f"完全排除同商户={exclude_all_same_merchant}"
        )

        pair_set, filter_stats, dense = self._recall_candidates(
            records, hashes, phash_threshold, dhash_threshold,
            filter_same_task, use_time_window, same_entity_window_days, exclude_all_same_merchant,
            max_candidates, dense_group_size
        )
        if self.is_cancelled:
            return {"status": "cancelled"}

        self.log_callback(f"初筛候选对总数：{len(pair_set):,} 对")
        if not pair_set:
            self.log_callback("未发现任何符合相似阈值与业务条件的疑似重复照片！", "WARN")
            return {"status": "empty", "candidates_count": 0}

        # 组装候选表数据
        self.status_callback("正在计算各区域汉明距离与业务判定...")
        candidates = self._build_candidate_frame(pair_set, hashes)
        self._add_metadata(candidates, records)
        self._add_business_context(candidates, records, same_entity_window_days if use_time_window else -1.0)

        # 优先级判定与聚类
        self.status_callback("正在执行聚类分组与优先级打分...")
        self._calc_priority(candidates)
        candidates = self._add_clusters(candidates)
        groups = self._group_summary(candidates)

        # 1. 导出相似清单（Excel & CSV）
        if export_list:
            self.status_callback("正在写入 Excel 相似清单...")
            self.log_callback(f"正在导出相似清单 Excel：{output_excel_path.name}")
            params = [
                ("工具版本", TOOL_VERSION),
                ("执行时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
                ("总照片数", total_photos),
                ("召回相似对数", len(candidates)),
                ("候选组聚类数", len(groups)),
                ("pHash阈值", phash_threshold),
                ("dHash阈值", dhash_threshold),
                ("同任务排除过滤", "已开启" if filter_same_task else "未开启"),
                ("时间窗口排除天数", f"{same_entity_window_days}天" if use_time_window else "关闭"),
                ("无条件排除同商户", "已开启" if exclude_all_same_merchant else "未开启"),
            ]
            for k, v in filter_stats.items():
                params.append((f"规则拦截-{k}", v))

            self._write_outputs(candidates, groups, dense, output_excel_path, output_csv_path, params, excel_limit)
            self.log_callback(f"✅ 相似清单已导出：{output_excel_path.name}", "SUCCESS")

        # 2. 导出带照片的离线复核包
        review_pkg_dir = None
        if export_review_package:
            self.status_callback("正在生成带照片的离线复核包...")
            self.log_callback("=== 开始打包带照片复核包 ===")
            pkg_root = output_excel_path.parent / f"review_packages_{datetime.now().strftime('%m%d_%H%M')}"
            review_pkg_dir = self.build_offline_review_package(
                candidates=candidates,
                output_dir=pkg_root,
                photo_root_dir=photo_root_dir,
                batch_size=100,
                priorities=["高", "中"],
                max_image_px=1600,
            )
            self.log_callback(f"✅ 离线网页复核包生成完成：{pkg_root.name}", "SUCCESS")

        # 3. 导出左右拼接并排对比图
        if export_pair_previews:
            self.status_callback("正在生成左右并排对比图...")
            preview_dir = output_excel_path.parent / f"相似照片对预览_{datetime.now().strftime('%m%d_%H%M')}"
            self._make_pair_previews(candidates, preview_dir, photo_root_dir, limit=200)
            self.log_callback(f"✅ 左右并排对比图已生成至：{preview_dir.name}", "SUCCESS")

        return {
            "status": "ok",
            "photos_count": total_photos,
            "candidates_count": len(candidates),
            "groups_count": len(groups),
            "excel_path": str(output_excel_path) if export_list else "",
            "csv_path": str(output_csv_path) if export_list else "",
            "review_pkg_dir": str(review_pkg_dir) if review_pkg_dir else "",
        }

    def _recall_candidates(
        self, records, hashes, phash_thresh, dhash_thresh,
        filter_same_task, use_time_window, window_days, exclude_all_same_merchant,
        max_candidates, dense_size
    ) -> Tuple[Set[int], dict, pd.DataFrame]:
        tasks = records.get("task_no", pd.Series([""] * len(records))).fillna("").astype(str).to_numpy()
        addresses = records.get("address_norm", pd.Series([""] * len(records))).fillna("").astype(str).to_numpy()
        merchants = records.get("merchant", pd.Series([""] * len(records))).map(normalize_merchant).to_numpy()
        submit_times = parse_submit_times(records)

        pair_set: Set[int] = set()
        stats = defaultdict(int)
        dense_rows = []
        fields = [(col, name, phash_thresh) for col, name in PHASH_COLUMNS]
        fields.append(("dhash", "dHash", dhash_thresh))

        for idx, (col, label, radius) in enumerate(fields, 1):
            if self.is_cancelled:
                break
            buckets: Dict[int, List[int]] = defaultdict(list)
            values = hashes[col]
            for i, v in enumerate(values):
                buckets[int(v)].append(i)

            for v, members in buckets.items():
                if len(members) >= dense_size:
                    dense_rows.append({
                        "哈希字段": label,
                        "哈希": f"{v:016x}",
                        "相同哈希照片数": len(members),
                        "任务数": len({tasks[i] for i in members if tasks[i]}),
                        "地址数": len({addresses[i] for i in members if addresses[i]}),
                        "示例文件": " | ".join(records.iloc[members[:5]]["filename"].astype(str)),
                    })

            unique_vals = list(buckets)
            self.log_callback(f"[{idx}/{len(fields)}] 构建 {label} BK-Tree：{len(unique_vals):,} 个唯一哈希")
            tree = BKTree(unique_vals)

            for step, val in enumerate(unique_vals, 1):
                if self.is_cancelled:
                    break
                near = tree.search(val, radius)
                first_members = buckets[val]
                for n_val in near:
                    if n_val < val:
                        continue
                    second_members = buckets[n_val]
                    if n_val == val:
                        for l_pos in range(len(first_members) - 1):
                            f = first_members[l_pos]
                            for r_pos in range(l_pos + 1, len(first_members)):
                                self._add_pair(
                                    pair_set, f, first_members[r_pos], tasks, addresses, merchants,
                                    submit_times, filter_same_task, use_time_window, window_days,
                                    exclude_all_same_merchant, stats, max_candidates
                                )
                    else:
                        for f in first_members:
                            for s in second_members:
                                self._add_pair(
                                    pair_set, f, s, tasks, addresses, merchants,
                                    submit_times, filter_same_task, use_time_window, window_days,
                                    exclude_all_same_merchant, stats, max_candidates
                                )
                if step % 5000 == 0 or step == len(unique_vals):
                    self.log_callback(f"  {label} 检索进度：{step:,}/{len(unique_vals):,}；当前候选对：{len(pair_set):,}")

        dense = pd.DataFrame(dense_rows)
        if not dense.empty:
            dense = dense.sort_values(["相同哈希照片数", "哈希字段"], ascending=[False, True])
        return pair_set, dict(stats), dense

    def _add_pair(
        self, pair_set, f, s, tasks, addresses, merchants, submit_times,
        filter_same_task, use_time_window, window_days, exclude_all_same_merchant,
        stats, max_c
    ):
        if f == s:
            return
        if f > s:
            f, s = s, f

        # 1. 过滤同任务号
        t1, t2 = tasks[f], tasks[s]
        if filter_same_task and t1 and t2 and t1 == t2:
            stats["同任务哈希命中"] += 1
            return

        # 2. 判断商户与地址
        a1, a2 = addresses[f], addresses[s]
        m1, m2 = merchants[f], merchants[s]
        same_addr = bool(a1 and a2 and a1 == a2)
        same_merch = bool(m1 and m2 and m1 == m2)

        # 若开启“无条件排除同商户”
        if exclude_all_same_merchant and same_merch:
            stats["无条件排除同商户命中"] += 1
            return

        # 3. 时间窗口过滤（仅同商户或同地址时生效）
        if (same_addr or same_merch) and use_time_window:
            tm1, tm2 = submit_times[f], submit_times[s]
            if np.isfinite(tm1) and np.isfinite(tm2):
                gap = abs(tm1 - tm2) / 86400.0
                if gap <= window_days:
                    # 在时间窗口内：正常复访，过滤！
                    if same_addr:
                        stats["窗口内同地址拦截"] += 1
                    if same_merch:
                        stats["窗口内同商户拦截"] += 1
                    return
                else:
                    # 超过时间窗口：跨月/跨期嫌疑，保留！
                    if same_addr:
                        stats["保留同地址跨期重复"] += 1
                    if same_merch:
                        stats["保留同商户跨期重复"] += 1
            else:
                stats["时间缺失同商户保留"] += 1

        key = (f << 32) | s
        pair_set.add(key)
        if len(pair_set) > max_c:
            raise RuntimeError(f"候选对超过安全上限 {max_c:,}！请调小阈值或检查密集哈希。")

    def _build_candidate_frame(self, pair_set: Set[int], hashes: Dict[str, np.ndarray]) -> pd.DataFrame:
        encoded = np.fromiter(pair_set, dtype=np.uint64, count=len(pair_set))
        left = (encoded >> np.uint64(32)).astype(np.int64)
        right = (encoded & np.uint64(0xFFFFFFFF)).astype(np.int64)
        frame = pd.DataFrame({"索引1": left, "索引2": right})

        dist_cols = []
        for col, label in PHASH_COLUMNS:
            out_col = f"pHash{label}距离"
            frame[out_col] = hamming_array(hashes[col], left, right)
            dist_cols.append(out_col)
        frame["dHash距离"] = hamming_array(hashes["dhash"], left, right)
        frame["aHash距离"] = hamming_array(hashes["ahash"], left, right)

        matrix = frame[dist_cols].to_numpy(dtype=np.uint8)
        best_pos = matrix.argmin(axis=1)
        frame["pHash最小距离"] = matrix[np.arange(len(frame)), best_pos]
        labels = np.array([name for _, name in PHASH_COLUMNS], dtype=object)
        frame["最佳pHash区域"] = labels[best_pos]
        frame = frame.sort_values(
            ["pHash最小距离", "dHash距离", "aHash距离", "索引1", "索引2"],
            ascending=True,
        ).reset_index(drop=True)
        return frame

    def _add_metadata(self, candidates: pd.DataFrame, records: pd.DataFrame) -> None:
        left = candidates["索引1"].astype(int).to_numpy()
        right = candidates["索引2"].astype(int).to_numpy()
        mapping = [
            ("filename", "文件名"), ("task_no", "task_no"),
            ("import_address", "import_address"), ("item_name", "item_name"),
            ("merchant_no", "merchant_no"), ("merchant", "merchant"),
            ("auditor_name", "auditor_name"), ("submit_time", "submit_time"),
            ("doc_url", "doc_url"), ("path", "完整路径"),
        ]
        for src, tgt in mapping:
            if src not in records.columns:
                vals = np.array([""] * len(records), dtype=object)
            else:
                vals = records[src].fillna("").astype(str).to_numpy()
            candidates[f"{tgt}1" if tgt in {"文件名", "完整路径"} else f"{tgt}_1"] = vals[left]
            candidates[f"{tgt}2" if tgt in {"文件名", "完整路径"} else f"{tgt}_2"] = vals[right]

    def _add_business_context(self, candidates: pd.DataFrame, records: pd.DataFrame, window_days: float) -> None:
        left = candidates["索引1"].astype(int).to_numpy()
        right = candidates["索引2"].astype(int).to_numpy()
        addresses = records.get("address_norm", pd.Series([""] * len(records))).fillna("").astype(str).to_numpy()
        merchants = records.get("merchant", pd.Series([""] * len(records))).map(normalize_merchant).to_numpy()
        submit_times = parse_submit_times(records)

        labels, gaps = [], []
        for f, s in zip(left, right):
            same_addr = bool(addresses[f] and addresses[s] and addresses[f] == addresses[s])
            same_merch = bool(merchants[f] and merchants[s] and merchants[f] == merchants[s])
            t1, t2 = submit_times[f], submit_times[s]
            valid_t = bool(np.isfinite(t1) and np.isfinite(t2))
            gap = abs(t1 - t2) / 86400.0 if valid_t else float("nan")
            gaps.append(round(gap, 4) if valid_t else np.nan)

            if same_merch and valid_t and window_days >= 0 and gap > window_days:
                labels.append("同商户跨期照片相似")
            elif same_addr and valid_t and window_days >= 0 and gap > window_days:
                labels.append("同地址跨期照片相似")
            elif (same_merch or same_addr) and not valid_t:
                labels.append("同商户或同地址但时间缺失")
            elif same_merch:
                labels.append("同商户照片相似")
            elif same_addr:
                labels.append("同地址照片相似")
            else:
                labels.append("跨商户照片相似")

        candidates["候选业务类型"] = labels
        candidates["时间差天"] = gaps

    def _calc_priority(self, candidates: pd.DataFrame) -> None:
        phash_min = candidates["pHash最小距离"].to_numpy(dtype=int)
        dhash_dist = candidates["dHash距离"].to_numpy(dtype=int)

        priority = np.array(["低"] * len(candidates), dtype=object)
        reason = np.array(["哈希距离接近候选阈值"] * len(candidates), dtype=object)

        mid = (phash_min <= 4) | ((phash_min <= 6) & (dhash_dist <= 6))
        priority[mid] = "中"
        reason[mid] = "四路 pHash 最小距离 <= 4 或组合哈希高相符"

        high = (phash_min <= 2) & (dhash_dist <= 4)
        priority[high] = "高"
        reason[high] = "极高置信度（最小 pHash <= 2 且 dHash <= 4）"

        candidates["复核优先级"] = priority
        candidates["判定依据"] = reason
        candidates["复核序号"] = np.arange(1, len(candidates) + 1)

    def _add_clusters(self, candidates: pd.DataFrame) -> pd.DataFrame:
        parent = {}
        def find(v):
            parent.setdefault(v, v)
            while parent[v] != v:
                parent[v] = parent[parent[v]]
                v = parent[v]
            return v
        def union(f, s):
            rf, rs = find(f), find(s)
            if rf != rs:
                parent[rs] = rf

        strong = candidates[candidates["复核优先级"].isin(["高", "中"])]
        for f, s in zip(strong["索引1"].astype(int), strong["索引2"].astype(int)):
            union(f, s)

        components = defaultdict(list)
        for v in list(parent):
            components[find(v)].append(v)
        valid = sorted((m for m in components.values() if len(m) >= 2), key=lambda x: (-len(x), min(x)))
        grp_map = {}
        for num, members in enumerate(valid, 1):
            grp = f"SIM-{num:05d}"
            for m in members:
                grp_map[m] = grp

        groups = []
        for f, s in zip(candidates["索引1"].astype(int), candidates["索引2"].astype(int)):
            gf, gs = grp_map.get(f, ""), grp_map.get(s, "")
            groups.append(gf if gf and gf == gs else "")
        candidates["候选组"] = groups
        return candidates

    def _group_summary(self, candidates: pd.DataFrame) -> pd.DataFrame:
        data = candidates[candidates["候选组"] != ""].copy()
        if data.empty:
            return pd.DataFrame(columns=["候选组", "照片数", "候选对数", "高优先级对数", "中优先级对数", "任务数"])
        rows = []
        for grp, frame in data.groupby("候选组"):
            files = set(frame["文件名1"]).union(set(frame["文件名2"]))
            tasks = set(frame["task_no_1"]).union(set(frame["task_no_2"])) - {""}
            rows.append({
                "候选组": grp,
                "照片数": len(files),
                "候选对数": len(frame),
                "高优先级对数": int((frame["复核优先级"] == "高").sum()),
                "中优先级对数": int((frame["复核优先级"] == "中").sum()),
                "任务数": len(tasks),
                "示例文件": " | ".join(sorted(files)[:5]),
            })
        return pd.DataFrame(rows).sort_values(["照片数", "候选对数"], ascending=False)

    def _write_outputs(self, candidates, groups, dense, out_excel, out_csv, params, excel_limit):
        out_excel.parent.mkdir(parents=True, exist_ok=True)
        out_csv.parent.mkdir(parents=True, exist_ok=True)

        preferred = [
            "候选组", "复核序号", "候选业务类型", "时间差天", "复核优先级", "判定依据",
            "文件名1", "文件名2", "task_no_1", "task_no_2",
            "import_address_1", "import_address_2", "item_name_1", "item_name_2",
            "merchant_no_1", "merchant_no_2", "merchant_1", "merchant_2",
            "auditor_name_1", "auditor_name_2", "submit_time_1", "submit_time_2",
            "doc_url_1", "doc_url_2", "完整路径1", "完整路径2",
            "最佳pHash区域", "pHash最小距离", "pHash原图距离",
            "pHash去底15%距离", "pHash去底25%距离", "pHash中心90%距离",
            "dHash距离", "aHash距离",
        ]
        existing = [c for c in preferred if c in candidates.columns]
        remainder = [c for c in candidates.columns if c not in existing and c not in {"索引1", "索引2"}]
        final_df = candidates[existing + remainder]

        final_df.to_csv(out_csv, index=False, encoding="utf-8-sig")

        limit = len(final_df) if excel_limit <= 0 else min(excel_limit, len(final_df))
        excel_candidates = final_df.head(limit)
        params_df = pd.DataFrame(params, columns=["参数", "值"])

        with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
            params_df.to_excel(writer, sheet_name="参数与汇总", index=False)
            excel_candidates.to_excel(writer, sheet_name="相似候选", index=False)
            groups.to_excel(writer, sheet_name="候选组汇总", index=False)
            dense.to_excel(writer, sheet_name="密集哈希监控", index=False)

        try:
            wb = load_workbook(out_excel)
            header_fill = PatternFill("solid", fgColor="D9EAF7")
            for sheet in wb.worksheets:
                sheet.freeze_panes = "A2"
                if sheet.max_row >= 1 and sheet.max_column >= 1:
                    sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
            wb.save(out_excel)
        except Exception:
            pass

    # ---------------- 导出带照片的离线复核包 ----------------
    def build_offline_review_package(
        self,
        candidates: pd.DataFrame,
        output_dir: Path,
        photo_root_dir: Optional[Path],
        batch_size: int = 100,
        priorities: Optional[List[str]] = None,
        max_image_px: int = 1600,
        jpeg_quality: int = 82,
    ) -> Path:
        output_dir.mkdir(parents=True, exist_ok=True)
        if priorities is None:
            priorities = ["高", "中"]

        filtered = candidates[candidates["复核优先级"].astype(str).isin(priorities)].copy()
        if filtered.empty:
            filtered = candidates.copy()

        batches = math.ceil(len(filtered) / batch_size)
        self.log_callback(f"纳入复核照片候选对数：{len(filtered):,}，共将拆分为 {batches} 个复核包")

        for b_idx in range(batches):
            if self.is_cancelled:
                break
            start = b_idx * batch_size
            b_data = filtered.iloc[start:start + batch_size]
            b_id = f"门头照复核_{b_idx + 1:03d}"
            pkg_dir = output_dir / b_id
            photos_dir = pkg_dir / "photos"
            photos_dir.mkdir(parents=True, exist_ok=True)

            records = []
            missing = []
            for _, row in b_data.iterrows():
                f1_path = self._locate_photo(row.get("完整路径1"), row.get("文件名1"), photo_root_dir)
                f2_path = self._locate_photo(row.get("完整路径2"), row.get("文件名2"), photo_root_dir)

                name1 = clean_text(row.get("文件名1"))
                name2 = clean_text(row.get("文件名2"))

                rel1 = ""
                rel2 = ""
                if f1_path and f1_path.is_file():
                    self._prepare_review_photo(f1_path, photos_dir / name1, max_image_px, jpeg_quality)
                    rel1 = f"photos/{name1}"
                else:
                    missing.append({"复核序号": row.get("复核序号"), "侧": 1, "文件名": name1})

                if f2_path and f2_path.is_file():
                    self._prepare_review_photo(f2_path, photos_dir / name2, max_image_px, jpeg_quality)
                    rel2 = f"photos/{name2}"
                else:
                    missing.append({"复核序号": row.get("复核序号"), "侧": 2, "文件名": name2})

                if rel1 or rel2:
                    records.append({
                        "id": f"PAIR-{clean_text(row.get('复核序号'))}",
                        "group": clean_text(row.get("候选组")) or f"PAIR-{clean_text(row.get('复核序号'))}",
                        "priority": clean_text(row.get("复核优先级")) or "中",
                        "reason": clean_text(row.get("判定依据")) or clean_text(row.get("候选业务类型")),
                        "image1": rel1,
                        "image2": rel2,
                        "file1": name1,
                        "file2": name2,
                        "task1": clean_text(row.get("task_no_1")),
                        "task2": clean_text(row.get("task_no_2")),
                        "address1": clean_text(row.get("import_address_1")),
                        "address2": clean_text(row.get("import_address_2")),
                        "item1": clean_text(row.get("item_name_1")),
                        "item2": clean_text(row.get("item_name_2")),
                        "auditor1": clean_text(row.get("auditor_name_1")),
                        "auditor2": clean_text(row.get("auditor_name_2")),
                        "time1": clean_text(row.get("submit_time_1")),
                        "time2": clean_text(row.get("submit_time_2")),
                        "phash": clean_text(row.get("pHash最小距离")),
                        "phashArea": clean_text(row.get("最佳pHash区域")),
                        "dhash": clean_text(row.get("dHash距离")),
                        "orbInliers": clean_text(row.get("ORB内点数")),
                        "orbRatio": clean_text(row.get("ORB内点比例")),
                    })

            # 写入 HTML 与使用说明
            data_json = json.dumps(records, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
            batch_json = json.dumps(b_id, ensure_ascii=False)
            html_content = HTML_TEMPLATE.replace("__BATCH_ID_JSON__", batch_json).replace("__CANDIDATES_JSON__", data_json)
            (pkg_dir / "开始复核.html").write_text(html_content, encoding="utf-8")
            (pkg_dir / "使用说明.txt").write_text(README_TEMPLATE, encoding="utf-8")
            if missing:
                pd.DataFrame(missing).to_csv(pkg_dir / "缺失照片.csv", index=False, encoding="utf-8-sig")

            self.log_callback(f"已生成复核包：{b_id} (包含 {len(records)} 对照片)")

        return output_dir

    def _locate_photo(self, full_path: Any, filename: Any, photo_root: Optional[Path]) -> Optional[Path]:
        if full_path:
            p = Path(str(full_path))
            if p.is_file():
                return p
        if photo_root and filename:
            name = str(filename)
            candidate = photo_root / name
            if candidate.is_file():
                return candidate
            matches = list(photo_root.rglob(name))
            if matches:
                return matches[0]
        return None

    def _prepare_review_photo(self, src: Path, dst: Path, max_px: int, quality: int):
        if dst.exists():
            return
        try:
            with Image.open(src) as img:
                fixed = ImageOps.exif_transpose(img).convert("RGB")
                w, h = fixed.size
                longest = max(w, h)
                if longest > max_px:
                    scale = max_px / float(longest)
                    nw, nh = max(1, int(w * scale)), max(1, int(h * scale))
                    fixed = fixed.resize((nw, nh), Image.Resampling.LANCZOS)
                fixed.save(dst, format="JPEG", quality=quality, optimize=True)
        except Exception:
            try:
                shutil.copy2(src, dst)
            except Exception:
                pass

    def _make_pair_previews(self, candidates: pd.DataFrame, preview_dir: Path, photo_root: Optional[Path], limit: int = 200):
        preview_dir.mkdir(parents=True, exist_ok=True)
        count = min(limit, len(candidates))
        for pos, (_, row) in enumerate(candidates.head(count).iterrows(), start=1):
            f1 = self._locate_photo(row.get("完整路径1"), row.get("文件名1"), photo_root)
            f2 = self._locate_photo(row.get("完整路径2"), row.get("文件名2"), photo_root)
            if not f1 or not f2:
                continue
            out_p = preview_dir / f"PAIR-{pos:04d}.jpg"
            if out_p.exists():
                continue
            try:
                with Image.open(f1) as img1, Image.open(f2) as img2:
                    im1 = ImageOps.exif_transpose(img1).convert("RGB")
                    im2 = ImageOps.exif_transpose(img2).convert("RGB")

                    # 统一缩放至高 600
                    h = 600
                    w1 = int(im1.width * (h / im1.height))
                    w2 = int(im2.width * (h / im2.height))
                    im1 = im1.resize((w1, h), Image.Resampling.LANCZOS)
                    im2 = im2.resize((w2, h), Image.Resampling.LANCZOS)

                    canvas = Image.new("RGB", (w1 + w2 + 10, h + 50), color=(30, 30, 30))
                    canvas.paste(im1, (0, 0))
                    canvas.paste(im2, (w1 + 10, 0))

                    draw = ImageDraw.Draw(canvas)
                    text = f"PAIR-{pos:04d} | {row.get('复核优先级')}优 | {row.get('候选业务类型')} | pHash={row.get('pHash最小距离')}"
                    draw.text((15, h + 15), text, fill=(255, 255, 255))
                    canvas.save(out_p, format="JPEG", quality=85)
            except Exception:
                pass


# ==============================================================================
# GUI 桌面界面系统
# ==============================================================================

class PhotoDuplicateScannerGUI:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("照片重复性智能检测与相似检索工作台 v2.1")
        self.root.geometry("1120x860")
        self.root.minsize(1000, 740)

        # 视觉风格
        self.bg_color = "#F5F6F8"
        self.card_bg = "#FFFFFF"
        self.primary_color = "#007AFF"
        self.success_color = "#34C759"
        self.warn_color = "#FF9500"
        self.danger_color = "#FF3B30"
        self.text_color = "#1D1D1F"
        self.text_secondary = "#86868B"

        self.root.configure(bg=self.bg_color)
        self.engine = DuplicateScanEngine(
            log_callback=self._log_msg,
            progress_callback=self._update_progress,
            status_callback=self._update_status,
        )
        self.is_running = False

        # 运行模式
        self.var_mode = tk.StringVar(value="auto")  # auto, recall_only, index_only

        # 数据路径
        self.var_photo_dir = tk.StringVar()
        self.var_mapping_file = tk.StringVar()
        self.var_sqlite_path = tk.StringVar()
        self.var_excel_path = tk.StringVar()
        self.var_csv_path = tk.StringVar()

        # 算法与过滤规则勾选项
        self.var_phash_thresh = tk.IntVar(value=8)
        self.var_dhash_thresh = tk.IntVar(value=8)
        self.var_filter_same_task = tk.BooleanVar(value=True)          # 过滤相同任务号
        self.var_use_time_window = tk.BooleanVar(value=True)           # 启用时间窗口
        self.var_window_days = tk.DoubleVar(value=7.0)                 # 时间窗口天数
        self.var_exclude_all_same_merchant = tk.BooleanVar(value=False)# 无条件排除同商户
        self.var_workers = tk.IntVar(value=max(2, min(8, (os.cpu_count() or 4) - 1)))
        self.var_refresh_cache = tk.BooleanVar(value=False)

        # 成果导出双轨勾选项
        self.var_export_list = tk.BooleanVar(value=True)               # 导出相似清单+原因 (Excel/CSV)
        self.var_export_review_pkg = tk.BooleanVar(value=False)        # 导出带照片离线复核包 (HTML+缩略图)
        self.var_export_previews = tk.BooleanVar(value=False)          # 导出并排对比图 (PAIR-xxxx.jpg)

        self._init_default_paths()
        self._build_ui()

    def _init_default_paths(self):
        cur_dir = Path(__file__).parent.resolve()
        sub_dir = cur_dir / "重复检查"

        # 1. 寻找现有 sqlite
        cand_dbs = [
            sub_dir / "photo_similarity_100k.sqlite",
            cur_dir / "photo_similarity_100k.sqlite",
            Path("/Volumes/t5/浙江区域采集照片/photo_similarity_100k.sqlite"),
        ]
        for db in cand_dbs:
            if db.is_file():
                self.var_sqlite_path.set(str(db))
                break
        if not self.var_sqlite_path.get():
            self.var_sqlite_path.set(str(cur_dir / "photo_similarity.sqlite"))

        # 2. 寻找映射表
        cand_maps = [
            cur_dir / "pos-jianhang2_已合并.xlsx",
            cur_dir / "pos-jianhang2.csv",
            sub_dir / "zhaopian.csv",
            cur_dir / "原始数据-绍兴.xlsx",
        ]
        for m in cand_maps:
            if m.is_file():
                self.var_mapping_file.set(str(m))
                break

        # 3. 寻找真正包含照片的目录
        cand_photos = [
            cur_dir / "pos-jianhang2_已合并_photos",
            Path("/Users/jun/Downloads/照片重复性检测/pos-jianhang2_已合并_photos"),
            Path("/Volumes/t5/浙江区域采集照片"),
            cur_dir / "门头照片",
            Path("/Users/jun/Desktop/照片重复性检测/pos-jianhang2_已合并_photos"),
        ]
        for p in cand_photos:
            if p.is_dir():
                try:
                    # 检查是否真正包含有效照片（跳过仅有 .part 的目录）
                    has_real_img = any(
                        f.is_file() and f.suffix.lower() in IMAGE_EXTENSIONS 
                        for f in list(p.iterdir())[:50]
                    )
                    if has_real_img:
                        self.var_photo_dir.set(str(p))
                        break
                except Exception:
                    pass
        if not self.var_photo_dir.get():
            for p in cand_photos:
                if p.is_dir():
                    self.var_photo_dir.set(str(p))
                    break

        # 4. 默认输出
        stamp = datetime.now().strftime("%m%d_%H%M")
        desktop = Path.home() / "Desktop"
        out_dir = desktop if desktop.is_dir() else cur_dir
        self.var_excel_path.set(str(out_dir / f"照片相似候选_{stamp}_7天窗口.xlsx"))
        self.var_csv_path.set(str(out_dir / f"照片相似候选_{stamp}_7天窗口.csv"))

    def _build_ui(self):
        top_bar = tk.Frame(self.root, bg=self.card_bg, height=65)
        top_bar.pack(fill="x", side="top")
        top_bar.pack_propagate(False)

        tk.Label(
            top_bar,
            text="🔍 照片重复性智能检测与相似检索工作台",
            font=("PingFang SC", 17, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
        ).pack(side="left", padx=25, pady=18)

        tk.Label(
            top_bar,
            text=f"v{TOOL_VERSION} ｜ 清单初筛 + 离线复核包双轨 ｜ 支持10万级秒级召回",
            font=("PingFang SC", 11),
            bg=self.card_bg,
            fg=self.text_secondary,
        ).pack(side="right", padx=25, pady=22)

        main_frame = tk.Frame(self.root, bg=self.bg_color)
        main_frame.pack(fill="both", expand=True, padx=20, pady=12)

        # ---------------- 卡片 1：运行模式选择 ----------------
        mode_card = tk.LabelFrame(
            main_frame,
            text=" 1. 运行模式选择 ",
            font=("PingFang SC", 12, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
            padx=15,
            pady=8,
        )
        mode_card.pack(fill="x", pady=(0, 8))

        modes_sub = tk.Frame(mode_card, bg=self.card_bg)
        modes_sub.pack(fill="x")

        tk.Radiobutton(
            modes_sub,
            text="🚀 一键全自动查重（新照片目录自动建库并查重）",
            variable=self.var_mode,
            value="auto",
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
            activebackground=self.card_bg,
            command=self._on_mode_changed,
        ).pack(side="left", padx=(0, 20))

        tk.Radiobutton(
            modes_sub,
            text="⚡ 仅相似检索（使用已有 SQLite 特征库，微调阈值秒级导出）",
            variable=self.var_mode,
            value="recall_only",
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
            activebackground=self.card_bg,
            command=self._on_mode_changed,
        ).pack(side="left", padx=(0, 20))

        tk.Radiobutton(
            modes_sub,
            text="📦 仅建库/更新（多进程提取特征存入 SQLite）",
            variable=self.var_mode,
            value="index_only",
            font=("PingFang SC", 11),
            bg=self.card_bg,
            activebackground=self.card_bg,
            command=self._on_mode_changed,
        ).pack(side="left")

        # ---------------- 卡片 2：路径与数据源配置 ----------------
        path_card = tk.LabelFrame(
            main_frame,
            text=" 2. 数据路径与导出设置 ",
            font=("PingFang SC", 12, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
            padx=15,
            pady=8,
        )
        path_card.pack(fill="x", pady=(0, 8))

        # 行 1：照片目录
        row1 = tk.Frame(path_card, bg=self.card_bg)
        row1.pack(fill="x", pady=3)
        tk.Label(row1, text="照片根目录:", width=13, anchor="e", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Entry(row1, textvariable=self.var_photo_dir, font=("PingFang SC", 11)).pack(side="left", fill="x", expand=True, padx=8)
        tk.Button(row1, text="选择目录...", command=self._choose_photo_dir, width=10).pack(side="left")

        # 行 2：映射 CSV/Excel
        row2 = tk.Frame(path_card, bg=self.card_bg)
        row2.pack(fill="x", pady=3)
        tk.Label(row2, text="任务映射表:", width=13, anchor="e", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Entry(row2, textvariable=self.var_mapping_file, font=("PingFang SC", 11)).pack(side="left", fill="x", expand=True, padx=8)
        tk.Button(row2, text="选择表格...", command=self._choose_mapping_file, width=10).pack(side="left")

        # 行 3：SQLite 库
        row3 = tk.Frame(path_card, bg=self.card_bg)
        row3.pack(fill="x", pady=3)
        tk.Label(row3, text="特征 SQLite 库:", width=13, anchor="e", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Entry(row3, textvariable=self.var_sqlite_path, font=("PingFang SC", 11)).pack(side="left", fill="x", expand=True, padx=8)
        tk.Button(row3, text="选择数据库...", command=self._choose_sqlite_path, width=10).pack(side="left")

        # 行 4：输出 Excel 路径
        row4 = tk.Frame(path_card, bg=self.card_bg)
        row4.pack(fill="x", pady=3)
        tk.Label(row4, text="导出候选清单:", width=13, anchor="e", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Entry(row4, textvariable=self.var_excel_path, font=("PingFang SC", 11)).pack(side="left", fill="x", expand=True, padx=8)
        tk.Button(row4, text="更改位置...", command=self._choose_excel_path, width=10).pack(side="left")

        # ---------------- 卡片 3：算法规则与业务过滤勾选 ----------------
        rule_card = tk.LabelFrame(
            main_frame,
            text=" 3. 算法阈值与业务过滤规则（支持自由勾选） ",
            font=("PingFang SC", 12, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
            padx=15,
            pady=8,
        )
        rule_card.pack(fill="x", pady=(0, 8))

        # 行 1：哈希阈值
        r_row1 = tk.Frame(rule_card, bg=self.card_bg)
        r_row1.pack(fill="x", pady=2)
        tk.Label(r_row1, text="四路 pHash 距离阈值:", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Spinbox(r_row1, from_=0, to=16, textvariable=self.var_phash_thresh, width=4, font=("PingFang SC", 11)).pack(side="left", padx=(4, 15))
        tk.Label(r_row1, text="差异 dHash 距离阈值:", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Spinbox(r_row1, from_=0, to=16, textvariable=self.var_dhash_thresh, width=4, font=("PingFang SC", 11)).pack(side="left", padx=(4, 15))
        tk.Label(r_row1, text="建库多进程数:", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left")
        tk.Spinbox(r_row1, from_=1, to=16, textvariable=self.var_workers, width=4, font=("PingFang SC", 11)).pack(side="left", padx=4)

        # 行 2：业务过滤勾选条
        r_row2 = tk.Frame(rule_card, bg=self.card_bg)
        r_row2.pack(fill="x", pady=4)

        tk.Checkbutton(
            r_row2,
            text="✔ 过滤相同任务号（排除同一任务内现场多张门头照）",
            variable=self.var_filter_same_task,
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
        ).pack(side="left", padx=(0, 15))

        tk.Checkbutton(
            r_row2,
            text="✔ 启用同商户/同地址时间窗口排除（排除短期正常复访）",
            variable=self.var_use_time_window,
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
        ).pack(side="left")

        tk.Label(r_row2, text="窗口天数:", font=("PingFang SC", 11), bg=self.card_bg).pack(side="left", padx=(5, 2))
        tk.Spinbox(r_row2, from_=0.0, to=365.0, increment=1.0, textvariable=self.var_window_days, width=5, font=("PingFang SC", 11)).pack(side="left", padx=2)
        tk.Label(r_row2, text="天（超过窗口视为跨月/跨期套用旧图保留！）", font=("PingFang SC", 10), fg=self.text_secondary, bg=self.card_bg).pack(side="left")

        # 行 3：其他高级过滤勾选
        r_row3 = tk.Frame(rule_card, bg=self.card_bg)
        r_row3.pack(fill="x", pady=2)

        tk.Checkbutton(
            r_row3,
            text="无条件排除所有同商户（仅看不同商户间的跨店套用作弊）",
            variable=self.var_exclude_all_same_merchant,
            font=("PingFang SC", 11),
            bg=self.card_bg,
        ).pack(side="left", padx=(0, 20))

        tk.Checkbutton(
            r_row3,
            text="强制刷新重新提取特征（忽略已有 SQLite 缓存）",
            variable=self.var_refresh_cache,
            font=("PingFang SC", 11),
            bg=self.card_bg,
        ).pack(side="left")

        # ---------------- 卡片 4：输出成果形式勾选 ----------------
        out_card = tk.LabelFrame(
            main_frame,
            text=" 4. 输出成果选择（先看清单，或同步带出照片复核包） ",
            font=("PingFang SC", 12, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
            padx=15,
            pady=8,
        )
        out_card.pack(fill="x", pady=(0, 8))

        out_sub = tk.Frame(out_card, bg=self.card_bg)
        out_sub.pack(fill="x")

        tk.Checkbutton(
            out_sub,
            text="📊 导出相似候选清单 + 判定原因 (Excel / CSV)  [推荐首选：秒级生成，快速总览]",
            variable=self.var_export_list,
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
        ).pack(anchor="w", pady=2)

        tk.Checkbutton(
            out_sub,
            text="🖼️ 同时导出带照片的离线复核包 (HTML 网页 + 本地照片缩略图)  [便于人工双图直观复核]",
            variable=self.var_export_review_pkg,
            font=("PingFang SC", 11, "bold"),
            bg=self.card_bg,
        ).pack(anchor="w", pady=2)

        tk.Checkbutton(
            out_sub,
            text="👥 同时生成左右并排对比拼接图 (PAIR-xxxx.jpg 图片文件夹)",
            variable=self.var_export_previews,
            font=("PingFang SC", 11),
            bg=self.card_bg,
        ).pack(anchor="w", pady=2)

        # ---------------- 卡片 5：运行控制与实时反馈 ----------------
        ctrl_card = tk.LabelFrame(
            main_frame,
            text=" 5. 执行控制与进度监控 ",
            font=("PingFang SC", 12, "bold"),
            bg=self.card_bg,
            fg=self.text_color,
            padx=15,
            pady=8,
        )
        ctrl_card.pack(fill="both", expand=True)

        btn_bar = tk.Frame(ctrl_card, bg=self.card_bg)
        btn_bar.pack(fill="x", pady=(0, 6))

        self.btn_run = tk.Button(
            btn_bar,
            text=" ▶ 开始执行检测与导出 ",
            font=("PingFang SC", 13, "bold"),
            bg=self.primary_color,
            fg="white",
            padx=20,
            pady=6,
            command=self._start_scan,
        )
        self.btn_run.pack(side="left", padx=(0, 10))

        self.btn_cancel = tk.Button(
            btn_bar,
            text=" ⏹ 中断停止 ",
            font=("PingFang SC", 11),
            state="disabled",
            padx=10,
            pady=6,
            command=self._cancel_scan,
        )
        self.btn_cancel.pack(side="left", padx=5)

        self.btn_open_excel = tk.Button(
            btn_bar,
            text=" 📊 打开生成的候选 Excel ",
            font=("PingFang SC", 11),
            state="disabled",
            padx=10,
            pady=6,
            command=self._open_result_excel,
        )
        self.btn_open_excel.pack(side="right", padx=5)

        self.btn_pack_from_list = tk.Button(
            btn_bar,
            text=" 🎯 从当前清单一键生成复核包 ",
            font=("PingFang SC", 11, "bold"),
            bg="#E8F4FD",
            padx=10,
            pady=6,
            command=self._pack_review_from_current_list,
        )
        self.btn_pack_from_list.pack(side="right", padx=5)

        # 进度条与状态
        p_frame = tk.Frame(ctrl_card, bg=self.card_bg)
        p_frame.pack(fill="x", pady=2)

        self.lbl_status = tk.Label(
            p_frame,
            text="就绪：请配置模式、勾选过滤规则与导出项后，点击「开始执行」",
            font=("PingFang SC", 11),
            fg=self.text_color,
            bg=self.card_bg,
            anchor="w",
        )
        self.lbl_status.pack(fill="x", side="top", pady=(0, 2))

        self.progressbar = ttk.Progressbar(p_frame, orient="horizontal", mode="determinate")
        self.progressbar.pack(fill="x", side="top")

        # 实时日志窗口
        self.txt_log = scrolledtext.ScrolledText(
            ctrl_card,
            height=7,
            font=("Menlo", 10),
            bg="#1E1E1E",
            fg="#D4D4D4",
            insertbackground="white",
        )
        self.txt_log.pack(fill="both", expand=True, pady=(6, 0))

        self.txt_log.tag_config("INFO", foreground="#CCCCCC")
        self.txt_log.tag_config("SUCCESS", foreground="#4EC9B0")
        self.txt_log.tag_config("WARN", foreground="#CE9178")
        self.txt_log.tag_config("ERROR", foreground="#F44747")

        self._on_mode_changed()

    def _on_mode_changed(self):
        mode = self.var_mode.get()
        if mode == "recall_only":
            self.lbl_status.config(text="【仅相似检索模式】：无需重新读图，基于已有 SQLite 特征库秒级完成过滤与导出！")
            self.btn_run.config(text=" ⚡ 开始快速相似检索与导出 ")
        elif mode == "index_only":
            self.lbl_status.config(text="【仅特征建库模式】：多进程扫描照片目录并提取指纹存入 SQLite（支持断点续跑）。")
            self.btn_run.config(text=" 📦 开始提取照片特征建库 ")
        else:
            self.lbl_status.config(text="【全自动查重模式】：自动先更新特征库，再基于 BK-Tree 检索并导出相似清单。")
            self.btn_run.config(text=" 🚀 开始全自动照片查重 ")

    # 路径选择器
    def _choose_photo_dir(self):
        p = filedialog.askdirectory(title="选择照片根目录", initialdir=self.var_photo_dir.get() or ".")
        if p:
            self.var_photo_dir.set(p)

    def _choose_mapping_file(self):
        p = filedialog.askopenfilename(
            title="选择任务映射表 (CSV 或 Excel)",
            filetypes=[("表格文件", "*.csv *.xlsx *.xls"), ("All Files", "*.*")],
            initialdir=Path(self.var_mapping_file.get()).parent if self.var_mapping_file.get() else ".",
        )
        if p:
            self.var_mapping_file.set(p)

    def _choose_sqlite_path(self):
        p = filedialog.askopenfilename(
            title="选择已有 SQLite 特征库或指定新库位置",
            filetypes=[("SQLite 数据库", "*.sqlite *.db"), ("All Files", "*.*")],
            initialdir=Path(self.var_sqlite_path.get()).parent if self.var_sqlite_path.get() else ".",
        )
        if p:
            self.var_sqlite_path.set(p)

    def _choose_excel_path(self):
        p = filedialog.asksaveasfilename(
            title="保存相似候选 Excel 结果",
            defaultextension=".xlsx",
            filetypes=[("Excel 表格", "*.xlsx")],
            initialdir=Path(self.var_excel_path.get()).parent if self.var_excel_path.get() else ".",
        )
        if p:
            self.var_excel_path.set(p)
            self.var_csv_path.set(str(Path(p).with_suffix(".csv")))

    # 日志与状态更新
    def _log_msg(self, msg: str, level: str = "INFO"):
        def append():
            timestamp = datetime.now().strftime("%H:%M:%S")
            self.txt_log.insert(tk.END, f"[{timestamp}] ", "INFO")
            self.txt_log.insert(tk.END, f"{msg}\n", level)
            self.txt_log.see(tk.END)
        self.root.after(0, append)

    def _update_status(self, text: str):
        self.root.after(0, lambda: self.lbl_status.config(text=text))

    def _update_progress(self, cur: int, total: int):
        def update():
            if total > 0:
                self.progressbar["maximum"] = total
                self.progressbar["value"] = cur
        self.root.after(0, update)

    # 启动与执行控制
    def _start_scan(self):
        if self.is_running:
            return

        mode = self.var_mode.get()
        db_path = Path(self.var_sqlite_path.get().strip()).expanduser().resolve()
        photo_dir = Path(self.var_photo_dir.get().strip()).expanduser().resolve() if self.var_photo_dir.get() else None
        mapping_file = Path(self.var_mapping_file.get().strip()).expanduser().resolve() if self.var_mapping_file.get() else None
        out_excel = Path(self.var_excel_path.get().strip()).expanduser().resolve()
        out_csv = Path(self.var_csv_path.get().strip()).expanduser().resolve()

        if mode in {"auto", "index_only"}:
            if not photo_dir or not photo_dir.is_dir():
                messagebox.showerror("错误", "请先选择有效的「照片根目录」！")
                return

        if mode == "recall_only":
            if not db_path.is_file():
                messagebox.showerror("错误", f"所选的特征 SQLite 数据库文件不存在：\n{db_path}\n\n请确认路径或切换为「全自动查重」模式。")
                return

        if not self.var_export_list.get() and not self.var_export_review_pkg.get() and not self.var_export_previews.get():
            messagebox.showwarning("提示", "请在「4. 输出成果选择」中至少勾选一项输出内容（如导出相似清单）！")
            return

        self.is_running = True
        self.engine.is_cancelled = False
        self.btn_run.config(state="disabled")
        self.btn_cancel.config(state="normal")
        self.btn_open_excel.config(state="disabled")
        self.txt_log.delete("1.0", tk.END)
        self.progressbar["value"] = 0

        threading.Thread(
            target=self._run_worker,
            args=(mode, photo_dir, mapping_file, db_path, out_excel, out_csv),
            daemon=True,
        ).start()

    def _cancel_scan(self):
        if messagebox.askyesno("确认", "确定要中断当前的查重检测吗？"):
            self.engine.cancel()
            self._log_msg("正在向引擎发送中断信号...", "WARN")

    def _run_worker(self, mode, photo_dir, mapping_file, db_path, out_excel, out_csv):
        start_time = time.time()
        try:
            # 阶段一：建库
            if mode in {"auto", "index_only"}:
                self._log_msg("=== 开始执行阶段一：提取照片特征与建库 ===", "INFO")
                res1 = self.engine.build_photo_index(
                    root_dir=photo_dir,
                    mapping_path=mapping_file,
                    db_path=db_path,
                    workers=self.var_workers.get(),
                    commit_every=200,
                    refresh=self.var_refresh_cache.get(),
                )
                if res1.get("status") == "cancelled":
                    self._on_finished("任务已被用户取消", success=False)
                    return

            # 阶段二：相似检索与双轨输出
            if mode in {"auto", "recall_only"}:
                self._log_msg("=== 开始执行阶段二：BK-Tree 空间检索与业务过滤 ===", "INFO")
                res2 = self.engine.scan_similarity_v3(
                    db_path=db_path,
                    output_excel_path=out_excel,
                    output_csv_path=out_csv,
                    phash_threshold=self.var_phash_thresh.get(),
                    dhash_threshold=self.var_dhash_thresh.get(),
                    filter_same_task=self.var_filter_same_task.get(),
                    use_time_window=self.var_use_time_window.get(),
                    same_entity_window_days=self.var_window_days.get(),
                    exclude_all_same_merchant=self.var_exclude_all_same_merchant.get(),
                    export_list=self.var_export_list.get(),
                    export_review_package=self.var_export_review_pkg.get(),
                    export_pair_previews=self.var_export_previews.get(),
                    photo_root_dir=photo_dir,
                )
                if res2.get("status") == "cancelled":
                    self._on_finished("任务已被用户取消", success=False)
                    return

                total_sec = time.time() - start_time
                msg = f"检测圆满完成！共召回 {res2.get('candidates_count', 0):,} 对疑似重复照片，耗时 {total_sec:.1f} 秒。"
                self._on_finished(msg, success=True, excel_path=str(out_excel) if self.var_export_list.get() else None)
            else:
                total_sec = time.time() - start_time
                self._on_finished(f"特征库构建完成，耗时 {total_sec:.1f} 秒！", success=True)

        except Exception as e:
            self._log_msg(f"运行发生异常：{e}", "ERROR")
            self._on_finished(f"执行失败：{e}", success=False)

    def _on_finished(self, msg: str, success: bool, excel_path: Optional[str] = None):
        def ui_done():
            self.is_running = False
            self.btn_run.config(state="normal")
            self.btn_cancel.config(state="disabled")
            if success:
                self.lbl_status.config(text=f"✅ {msg}")
                self._log_msg(msg, "SUCCESS")
                if excel_path and Path(excel_path).is_file():
                    self.btn_open_excel.config(state="normal")
                    if messagebox.askyesno("检测完成", f"{msg}\n\n是否立即打开导出的候选清单 Excel 表格？"):
                        self._open_result_excel()
            else:
                self.lbl_status.config(text=f"❌ {msg}")
                messagebox.showerror("提示", msg)
        self.root.after(0, ui_done)

    def _open_result_excel(self):
        p = self.var_excel_path.get()
        if p and Path(p).is_file():
            subprocess.run(["open", p])
        else:
            messagebox.showwarning("提示", "候选 Excel 文件尚未生成或已被移除。")

    # 从当前清单快速打包为复核包
    def _pack_review_from_current_list(self):
        excel_path = Path(self.var_excel_path.get().strip()).expanduser().resolve()
        if not excel_path.is_file():
            # 提示用户浏览已有的 Excel 清单
            chosen = filedialog.askopenfilename(
                title="选择已生成的相似候选 Excel 清单",
                filetypes=[("Excel 表格", "*.xlsx"), ("All Files", "*.*")],
                initialdir=excel_path.parent if excel_path.parent.is_dir() else ".",
            )
            if not chosen:
                return
            excel_path = Path(chosen)

        photo_root = Path(self.var_photo_dir.get().strip()).expanduser().resolve() if self.var_photo_dir.get() else None
        if not photo_root or not photo_root.is_dir():
            chosen_dir = filedialog.askdirectory(title="选择照片根目录（用于抽取并压缩照片）")
            if not chosen_dir:
                return
            photo_root = Path(chosen_dir)

        try:
            self._log_msg(f"正在从清单读取数据：{excel_path.name}")
            data = pd.read_excel(excel_path, sheet_name="相似候选")
            pkg_root = excel_path.parent / f"review_packages_{datetime.now().strftime('%m%d_%H%M')}"
            self._log_msg(f"开始打包复核包至：{pkg_root.name} ...")
            self.engine.build_offline_review_package(
                candidates=data,
                output_dir=pkg_root,
                photo_root_dir=photo_root,
                batch_size=100,
                priorities=["高", "中"],
                max_image_px=1600,
            )
            self._log_msg(f"✅ 复核包打包圆满完成：{pkg_root}", "SUCCESS")
            if messagebox.askyesno("打包完成", f"已成功从清单生成带照片的离线复核包！\n路径：{pkg_root}\n\n是否立即在访达中打开？"):
                subprocess.run(["open", str(pkg_root)])
        except Exception as e:
            messagebox.showerror("打包失败", f"从清单生成复核包失败：{e}")


def main():
    root = tk.Tk()
    app = PhotoDuplicateScannerGUI(root)
    root.mainloop()


if __name__ == "__main__":
    main()
